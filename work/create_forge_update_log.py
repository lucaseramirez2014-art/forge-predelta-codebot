from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path("/Users/lucasramirez/Documents/Codex/2026-10-01/ere-is-my-current-app-dashboard/outputs/FORGE_Version_0_1_Beta_Update_Log.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=10.5, bold=False, color="1F1F1F"):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", style=None, space_before=0, space_after=7, size=10.5, color="292929", bold=False):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.18
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(15)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run, size=14, bold=True, color="000000")
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "2C2C2C")
        set_cell_border(cell, color="BFBFBF")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_font(run, size=9.5, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, "F5F5F5" if row_index % 2 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            run = p.add_run(value)
            set_font(run, size=9.5, color="242424")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)

header = section.header
header_p = header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_p.paragraph_format.space_after = Pt(0)
run = header_p.add_run("FORGE PRODUCT UPDATE")
set_font(run, size=8, bold=True, color="5C5C5C")

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = Pt(5)
run = footer_p.add_run("FORGE Version 0 1 Beta")
set_font(run, size=8, color="777777")

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(5)
title_run = title.add_run("FORGE Version 0 1 Beta Update Log")
set_font(title_run, size=25, bold=True, color="000000")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(15)
subtitle_run = subtitle.add_run("Implementation summary and handoff notes")
set_font(subtitle_run, size=11, color="5C5C5C")

add_paragraph(
    doc,
    "FORGE v0.1 Beta is a complete local Next.js prototype for gamified productivity and skill development. The beta implements the planned loop from campaign selection through focused work, quest completion, XP gain, and dashboard progress. It uses mock data only and keeps the codebase ready for future persistence, accounts, AI generation, and native mobile work.",
    space_after=11,
)

add_heading(doc, "Current Scope")
add_table(
    doc,
    ["Route", "Purpose"],
    [
        ("/", "Product entry page with the beta concept and campaign entry points."),
        ("/dashboard", "Shows level, XP progress, streak, active campaign, next quest, and quest board."),
        ("/campaign", "Shows the Python Developer campaign, available quests, progress, and the upcoming boss quest."),
        ("/focus", "Runs a quest-specific focus session with a countdown timer and a completion action."),
        ("/complete", "Shows quest completion, XP earned, and links back to progress or the campaign."),
    ],
    [1.2, 5.55],
)

add_heading(doc, "Core Product Loop")
add_paragraph(doc, "A user opens the dashboard, selects a quest, and enters a focus session. Completing the session records the quest in the in-memory beta state, adds the quest reward to XP, and presents the completion screen. Returning to the dashboard shows the updated XP total, cleared quest count, campaign progress, and completed quest state.")

add_heading(doc, "Reusable Components")
add_table(
    doc,
    ["Component", "Role"],
    [
        ("XPBar", "Displays current XP against the next level target."),
        ("QuestCard", "Presents difficulty, focus duration, XP reward, completion state, and start action."),
        ("CampaignCard", "Shows the active campaign, completion progress, and campaign entry point."),
        ("StatCard", "Displays streak, focus time, and cleared quest statistics."),
        ("FocusTimer", "Runs the focus countdown with pause, reset, and completion behavior."),
    ],
    [1.6, 5.15],
)

add_heading(doc, "Mock Data and Boundaries")
add_paragraph(doc, "The beta includes a Python Developer campaign with three available practice quests and one upcoming boss quest. Mock data defines the user profile, level, XP, streak, campaign metadata, quest difficulty, reward, duration, and focus prompt. Progress exists only for the current browser session. The project does not include authentication, a database, AI APIs, payments, social features, leaderboards, or screen-time blocking.")

add_heading(doc, "Project Structure")
add_table(
    doc,
    ["Location", "Contents"],
    [
        ("app", "App Router pages, root layout, and global styling."),
        ("components", "Reusable UI cards, navigation, timer, XP bar, and in-memory progress provider."),
        ("lib", "Mock user, campaign, quest data, and shared TypeScript types."),
    ],
    [1.6, 5.15],
)

add_heading(doc, "Verification")
add_paragraph(doc, "The project passed ESLint and a production Next.js build. All requested routes are statically generated successfully. The implementation uses standard Next.js, React, TypeScript, and Tailwind packages with no additional runtime dependencies.")

add_heading(doc, "Recommended Next Work")
add_paragraph(doc, "Future work can replace the in-memory provider with persistent user progress, add campaign creation, and connect account or AI features without changing the current route and component boundaries. Those features are intentionally outside the v0.1 Beta scope.")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
